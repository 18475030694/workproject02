# -*- coding: utf-8 -*-
"""新秘宝 E2E 结果回填：data.json → 用例 Excel 加【执行结果/复判/执行说明/截图】4列 + 概要Sheet；
报告 MD 落 temp_case/reports/。用法: python backfill_mibao_e2e.py [data_json路径]"""
import json
import sys
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

ROOT = Path(r"E:\workproject")
CASES_XLSX = ROOT / "temp_case" / "cases" / "新秘宝用例_v2执行版.xlsx"
REPORT_DIR = ROOT / "temp_case" / "reports"

data_path = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "reports" / "e2e_mibao_latest" / "data.json"
data = json.loads(data_path.read_text(encoding="utf-8"))
res = {}
for c in data.get("cases", []):
    cid = c["case_id"]
    if cid in res:  # 同一用例多条记录（如 M1 逐档位、C1 多次采样）→ 合并备注
        prev = res[cid]
        prev["note"] = " | ".join(dict.fromkeys(
            x for x in [str(prev.get("note") or ""), str(c.get("note") or "")] if x))
        if prev.get("passed") is None and c.get("passed") is not None:
            prev["passed"] = c["passed"]
        continue
    res[cid] = c

wb = openpyxl.load_workbook(CASES_XLSX)
ws = wb["测试用例清单"]
COLS = ["执行结果", "复判", "执行说明", "截图"]
hdr = ws.max_column
for j, h in enumerate(COLS, hdr + 1):
    c = ws.cell(row=1, column=j, value=h)
    c.font = Font(bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor="548235")
fills = {"PASS": PatternFill("solid", fgColor="C6EFCE"),
         "FAIL": PatternFill("solid", fgColor="FFC7CE"),
         "NA": PatternFill("solid", fgColor="D9D9D9")}
stat = {"PASS": 0, "FAIL": 0, "NA": 0, "NA_未覆盖": 0}
matched = set()
for i in range(2, ws.max_row + 1):
    cid = ws.cell(row=i, column=2).value
    e = res.get(cid)
    if not e:
        # data.json 无记录：脚本未覆盖/合并至相邻用例 → 标 NA 未覆盖
        stat["NA_未覆盖"] += 1
        ws.cell(row=i, column=hdr + 1, value="NA").fill = fills["NA"]
        ws.cell(row=i, column=hdr + 1).alignment = Alignment(horizontal="center")
        ws.cell(row=i, column=hdr + 2, value="—")
        ws.cell(row=i, column=hdr + 3, value="本轮脚本未覆盖（合并至相邻用例或采样条件未满足），转人工/下轮补测")
        continue
    matched.add(cid)
    tag = "PASS" if e["passed"] is True else ("FAIL" if e["passed"] is False else "NA")
    stat[tag] += 1
    note = str(e.get("note") or "")
    shot = str(e.get("screenshot") or "")
    ws.cell(row=i, column=hdr + 1, value=tag).fill = fills[tag]
    ws.cell(row=i, column=hdr + 1).alignment = Alignment(horizontal="center")
    ws.cell(row=i, column=hdr + 2, value="待复判" if tag == "FAIL" else ("—" if tag == "NA" else "通过"))
    ws.cell(row=i, column=hdr + 3, value=note[:500])
    ws.cell(row=i, column=hdr + 4, value=shot)
for j in range(hdr + 1, hdr + 5):
    ws.column_dimensions[openpyxl.utils.get_column_letter(j)].width = {"执行结果": 9, "复判": 8, "执行说明": 46, "截图": 24}[ws.cell(row=1, column=j).value]

# 概要
AUX_IDS = {"PREP", "LOGIN", "OPENUI", "MIB-F2b", "MIB-G4b", "MIB-K1b", "MIB-D4"}
aux_pass = sorted(c["case_id"] for c in data.get("cases", [])
                  if c["case_id"] in AUX_IDS and c["passed"] is True)
ws2 = wb.create_sheet("E2E概要")
ws2.append(["新秘宝(仙机) E2E 执行概要"])
ws2["A1"].font = Font(bold=True, size=13)
rows = [
    ("执行时间", data.get("start_time")), ("账号", f"{data.get('account')}@{data.get('server')}"),
    ("产物", str(data.get("out_dir"))), ("PASS", stat["PASS"]), ("FAIL", stat["FAIL"]),
    ("NA(人工/缓测)", stat["NA"]), ("NA(脚本未覆盖→人工)", stat["NA_未覆盖"]),
    ("辅助步骤(全PASS)", ", ".join(aux_pass)),
]
for r in rows:
    ws2.append(r)
ws2.append([])
ws2.append(["FAIL/NA 明细"])
ws2.cell(row=ws2.max_row, column=1).font = Font(bold=True)
for c in data.get("cases", []):
    if c["passed"] is not True:
        ws2.append([c["case_id"], c["title"], "FAIL" if c["passed"] is False else "NA", str(c.get("note") or "")[:200]])

out_x = CASES_XLSX.with_name(CASES_XLSX.stem + "_已回填.xlsx")
wb.save(out_x)

# MD 报告
fails = [c for c in data.get("cases", []) if c["passed"] is False]
nas = [c for c in data.get("cases", []) if c["passed"] is None]
L = ["# 新秘宝(仙机) E2E 执行报告", "",
     f"- 时间: {data.get('start_time')}  账号: {data.get('account')}@{data.get('server')}",
     f"- 汇总: **PASS {stat['PASS']} / FAIL {stat['FAIL']} / NA {stat['NA']+stat['NA_未覆盖']}**（NA=人工/缓测 {stat['NA']} + 脚本未覆盖 {stat['NA_未覆盖']}）",
     f"- 产物: {data.get('out_dir')}", "",
     "## 关键静态发现（需求 vs 实现）", "",
     "1. 倍率：服务端按权重随机 {1:50%,2:37%,4:9%,10:4%} 整批生效，与需求“1~4倍图案组合”不一致（MIB-E6 待策划确认）。",
     "2. 3倍图资源存在但配置不可达（treasurebox_img03_3beijiangli）。",
     "3. 货币即“仙机”（道具1001~1005），不足时按钮置灰“暂无仙机”，协议兜底 1027。",
     "4. 默认档位=持有数在[1,5,10,50,200]中≤持有数的最大档。", ""]
if fails:
    L += ["## FAIL 明细", "", "| 用例 | 标题 | 说明 |", "|---|---|---|"]
    L += [f"| {c['case_id']} | {c['title']} | {str(c.get('note') or '')[:180]} |" for c in fails]
    L.append("")
L += ["## NA 清单", "", "| 用例 | 标题 | 原因 |", "|---|---|---|"]
L += [f"| {c['case_id']} | {c['title']} | {str(c.get('note') or '')[:120]} |" for c in nas]
L += ["", "## 本轮脚本未覆盖（转人工/下轮补测）", "", "| 用例 |", "|---|"]
# 用 Excel 行取未覆盖 ID 标题
for i in range(2, ws.max_row + 1):
    cid = ws.cell(row=i, column=2).value
    if cid in res:
        continue
    L.append(f"| {cid} {ws.cell(row=i, column=4).value} |")
REPORT_DIR.mkdir(parents=True, exist_ok=True)
out_m = REPORT_DIR / "新秘宝E2E执行报告.md"
out_m.write_text("\n".join(L), encoding="utf-8")
print("backfill:", out_x)
print("report :", out_m)
print("summary:", stat)
