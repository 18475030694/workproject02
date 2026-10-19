# -*- coding: utf-8 -*-
"""校验回填文件：新sheet结构、TC-041/TC-075落位"""
import openpyxl

wb = openpyxl.load_workbook(r'E:\workproject\temp_case\cases\基金特权入口合并测试用例_已回填_v2.xlsx')
print('sheets:', wb.sheetnames)
ws = wb['非通过用例汇总']
print('rows:', ws.max_row - 1, 'cols:', ws.max_column)
print('headers:', [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)])

ids = [str(ws.cell(row=r, column=1).value) for r in range(2, ws.max_row + 1)]
verdicts = {}
for r in range(2, ws.max_row + 1):
    v = ws.cell(row=r, column=10).value
    verdicts[v] = verdicts.get(v, 0) + 1
print('verdict distribution:', verdicts)
print('TC-041 in sheet:', 'TC-041' in ids, '| TC-075 in sheet:', 'TC-075' in ids)
# 备注列应全空
remark_empty = all(ws.cell(row=r, column=11).value in (None, '') for r in range(2, ws.max_row + 1))
print('备注列全空:', remark_empty)

# 主表中确认两条用例的最终结论
ws1 = wb['基金特权入口合并测试用例']
for r in range(2, ws1.max_row + 1):
    tc = str(ws1.cell(row=r, column=1).value or '').strip()
    if tc in ('TC-041', 'TC-075'):
        print(tc, '| verdict:', ws1.cell(row=r, column=10).value)
