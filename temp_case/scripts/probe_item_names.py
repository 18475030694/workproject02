# -*- coding: utf-8 -*-
import io, sys
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
import openpyxl
wb = openpyxl.load_workbook(r'E:\p1-x\trunk\export\D-道具.xlsx', read_only=True, data_only=True)
ws = wb[wb.sheetnames[0]]
rows = list(ws.iter_rows(values_only=True))
hdr_i = None
for i, r in enumerate(rows[:8]):
    if r and r[0] == '导出字段':
        hdr_i = i
        print('fields:', r[:10])
        break
for r in rows[hdr_i + 1:]:
    if r and r[1] in (5, 30005, 1000, 1006, 1007):
        print(r[1], '|', r[2])
