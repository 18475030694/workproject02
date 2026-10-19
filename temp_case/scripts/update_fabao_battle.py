# -*- coding: utf-8 -*-
"""法宝战斗表现回归结论回填（trae · 2026-09-09）——更新 法宝UI优化测试用例_trae_v2.xlsx
依据：battle1(13:02)/battle2(13:22)/battle3(13:28) 三轮战斗回归取证
- battle2: GMRoot入口自检(GMBTN/BATTLE_TEST/MAIN_BATTLE_TEST均存在且激活)、TC063 GMBTN toggle打通
- battle3: B010实锤账号无英雄(编队+heroMap均空)、UITestFight打开+BtnFight_可点、战斗阶段0新增报错、无灵力/年份残留"""
import openpyxl
from openpyxl.styles import PatternFill, Alignment

DST = r"E:\workproject\temp_case\cases\法宝UI优化测试用例_trae_v2.xlsx"
OUT = DST

# 战斗回归新增证据：{TC: (新结论, 结果标签)}
ROUND_BATTLE = {
    "TC063": ("PASS(战斗回归battle2/battle3运行时)：GM面板真实入口=GMRoot/GM下的GMBTN按钮(UIMgr.ts L261-265)，clickNode('UIRootCamera/GMRoot/GM/GMBTN')打开后Global.uiMgr.getUI('UIGM').visible()=True、再点关闭后visible=False，开/关toggle验证通过；此前两轮FAIL系UI注册名'UIGM'与面板节点名'uiGM'不一致导致findByName未命中（工具误判，非产品缺陷）", "通过"),
    "TC065": ("待人工复核(战斗回归battle3部分验证)：战斗测试入口BATTLE_TEST(UIMgr.ts L337-340)点击生效、UITestFight战斗界面正常打开、『战斗』按钮(BtnFight_)可点击，战斗运行阶段0新增JS报错、界面文本无『灵力/年份』残留；但战斗内法宝技能释放需『已装配法宝+编队英雄』，测试账号1353当前编队与heroMap均为空（无英雄可装配，battle1-3三轮实锤）→技能特效/数值表现建议人工用正常账号进入战斗复核", "待人工复核"),
    "TC024": ("PASS(静态+运行时，含战斗场景)：soulring业务代码已无灵力文案(formatYear成死代码，probe2实锤仅死路径输出灵力)；v1/v2运行时主城/图鉴/炼宝均未见『灵力/万灵力』；battle3战斗测试界面文本快照(t1/t2)亦无『灵力/年份』命中，全局覆盖完成", "通过"),
}

STATUS = {
    "通过": "FFC6EFCE", "部分实现": "FFFFEB9C", "未实现": "FFFFC7CE", "失败": "FFFFC7CE",
    "待运行时": "FFD9D9D9", "资源就绪": "FFFFEB9C", "静态关注": "FFFFEB9C",
    "待人工复核": "FFDDEBF7",
}

wb = openpyxl.load_workbook(DST)
ws = wb["测试用例"]

hdr_row = 1
id_col, conc_col, res_col = 1, None, None
for c in range(1, ws.max_column + 1):
    v = ws.cell(row=hdr_row, column=c).value
    if v == "用例编号":
        id_col = c
    if v and str(v).startswith("复测结论"):
        conc_col = c
    if v == "结果":
        res_col = c
assert conc_col and res_col, f"列未找到 conc={conc_col} res={res_col}"

n_upd = 0
for row in range(hdr_row + 1, ws.max_row + 1):
    cid = ws.cell(row=row, column=id_col).value
    if not cid:
        continue
    cid = str(cid).strip()
    if cid not in ROUND_BATTLE:
        continue
    conc, label = ROUND_BATTLE[cid]
    c1 = ws.cell(row=row, column=conc_col, value=conc)
    c1.alignment = Alignment(wrap_text=True, vertical="top")
    c2 = ws.cell(row=row, column=res_col, value=label)
    c2.fill = PatternFill("solid", fgColor=STATUS.get(label, "FFFFFFFF"))
    c2.alignment = Alignment(vertical="center", horizontal="center")
    n_upd += 1

try:
    wb.save(OUT)
except PermissionError:
    OUT = r"E:\workproject\temp_case\cases\法宝UI优化测试用例_trae_v3.xlsx"
    wb.save(OUT)
print("已更新:", OUT)
print("更新用例数:", n_upd)

# 统计最终分布
from collections import Counter
stat = Counter()
for row in range(hdr_row + 1, ws.max_row + 1):
    v = ws.cell(row=row, column=res_col).value
    if v:
        stat[str(v).strip()] += 1
total = sum(stat.values())
print("最终分布:", dict(stat))
print("总计:", total)
for k, n in stat.most_common():
    print(f"  {k}: {n} ({n/total*100:.0f}%)")
