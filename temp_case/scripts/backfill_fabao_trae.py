# -*- coding: utf-8 -*-
"""法宝UI优化 全量重测结论回填（trae · 2026-09-08）——生成 法宝UI优化测试用例_trae.xlsx
依据：静态源码/配置核对（权威，含 9/8 17:10 最新改动）+ 运行时三轮取证（v1 17:18 / v2 16:07 / probe 17:26）
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

SRC = r"E:\workproject\temp_case\cases\法宝UI优化测试用例.xlsx"
DST = r"E:\workproject\temp_case\cases\法宝UI优化测试用例_trae.xlsx"

# 每条用例的复测结论（静态权威 + 运行时佐证；已并入 17:10 新改动与 probe 证据）
CONC = {
    "TC001": "部分实现：资源已数据化到cfg_soul_ring_quality(10档bg)，但代码仍按quality取道具框(UISoulRing.ts L51-52/83-84)；配置quality==ring_group 100%，显示巧合等价",
    "TC002": "部分实现：光球draw_ball_eff_res已配10档，但代码按quality取；同上巧合等价",
    "TC003": "部分实现：漩涡eff_res已配10档，但代码按quality取(frame_eff_res来自qualityCfg)",
    "TC004": "部分实现：技能框skill_res已配10档，但取用仍按quality",
    "TC005": "部分实现：图鉴框atlas_res已配10档，但取用仍按quality",
    "TC006": "未实现：名称色值仍按quality(UIHeroSoulItem L96等 cfg_soul_ring_quality[quality].color)，未按ring_group；需求色值本身标『等美术给』",
    "TC007": "无法验证：配置quality==ring_group 100%，需GM造quality≠ring_group数据才能暴露",
    "TC008": "无法验证：同上，当前数据无法区分",
    "TC009": "无法验证：代码读quality，改quality必然影响外观",
    "TC010": "PASS(静态)：10组资源文件齐全(bg/光球/漩涡/技能框/图鉴框/筛选图均存在)",
    "TC011": "待运行时：cfg同时保留quality与ring_group字段，数据兼容；运行时受预览新旧混杂阻塞(D1)",
    "TC012": "PASS(静态+运行时)：UISoulRing列表项已用formatLevel输出阶；probe2运行时实锤 formatLevel(50,6001)→『50阶』",
    "TC013": "PASS(静态)：UISoulRingShow已用formatLevel",
    "TC014": "PASS(静态+运行时佐证)：UISoulRingAtlasDesc L52已改『初始阶数』+formatLevel；probe2实锤formatLevel→『X阶』；但预览图鉴详情旧chunk仍显『初始阶数：1000年』(D1)",
    "TC015": "PASS(静态)：UIRingUnlockSkill L61用formatLevel输出阶",
    "TC016": "PASS(静态)：UIRingUnlockFeatures L59用formatLevel",
    "TC017": "PASS(静态)：UISoulRingDrawRewards/UIRingTipsTemp已用formatLevel",
    "TC018": "PASS(静态+运行时)：抽奖结果列表已改阶(UIHeroSoulItem/UIHeroSoulRing L526/563)；v1运行时抽奖页『新加入奖池』列表实显1阶/5阶",
    "TC019": "PASS(静态)：UISoulRingUpgradeResult L112用formatLevel",
    "TC020": "PASS(静态)：RingAtlasItem L85用formatLevel；运行时图鉴item仍显『80万年』(预览旧chunk，D1)",
    "TC021": "PASS(静态)：推荐条目同列表项链路，已改阶",
    "TC022": "PASS(静态)：UIHeroSoulItem L90/169、UIHeroSoulRing L526已改阶(17:10版Year_rich=formatLevel)",
    "TC023": "部分实现：SoulRingCtl L69仍用formatYear2(图标档位索引，非文案)，战斗文案待运行时",
    "TC024": "PASS(静态+运行时)：soulring业务代码已无灵力文案(formatYear成死代码，probe2实锤formatYear→『1000灵力』仅死路径)；v1/v2运行时均未见『灵力/万灵力』",
    "TC025": "未实现：阶数颜色仍按quality(UIHeroSoulItem L89-90)或灰#8999a1(formatLevel默认)，未按level取色",
    "TC026": "PASS(静态)：数据层已按ring_group定位升级表(SoulRingMgr L369/372)，阶数与加成计算一致",
    "TC027": "PASS(静态)：解锁条件文案已随formatLevel改阶(UIRingUnlockSkill/UIRingUnlockFeatures)",
    "TC028": "待运行时：达到阶数即时解锁依赖服务端事件，运行时受预览混杂阻塞(D1)",
    "TC029": "未实现：17:10版UIHeroSoulItem L156-179仍显示全部特性(灰色显实际效果文案)，无『未解锁/要求阶数』文案；v2运行时取证一致",
    "TC030": "未实现：无『暂无特性』文案与判定逻辑",
    "TC031": "未实现：特性显示逻辑未按需求实现，边界无法正确验证",
    "TC032": "未实现：同上",
    "TC033": "未实现：各界面特性显示规则不一致(UISoulRingDesc显示全部/UIHeroSoulItem显示全部/UIRingUnlockFeatures只显示已激活)",
    "TC034": "未实现：UISoulRingDesc L123-151显示全部特性(灰/亮区分)，未『只显示已激活』",
    "TC035": "未实现：无『无已激活特性』特殊表现逻辑",
    "TC036": "待运行时：全部激活时显示完整，运行时受预览混杂阻塞(D1)",
    "TC037": "未实现：cfg_soul_ring_pool_lv_desc仍year(品质名)+color色值；UISoulRingDrawLvDesc L121/146仍用e.year+e.color；v1/v2运行时均实显『精品/良品』",
    "TC038": "未实现：未新增组别底图资源字段",
    "TC039": "未实现：色值配置未废弃仍生效",
    "TC040": "PASS(运行时)：升级预览LV.1→LV.2概率数值链路正常(cfg_soul_ring_pool_lv_desc)；v1/v2运行时弹窗均可打开",
    "TC041": "部分实现：法宝列表项UISoulRing无攻防血；装配item UIHeroSoulItem L99-110有基础属性显示(17:10版仍保留)",
    "TC042": "待运行时：属性与角色面板一致性需运行时比对，受预览混杂阻塞(D1)",
    "TC043": "待运行时：名称上移/长文本遮挡需运行时，受预览混杂阻塞(D1)",
    "TC044": "部分实现：分类仍按quality(UISoulRingAtlas qualityList来自cfg_soul_ring_quality)，未按ring_group；v1运行时图鉴按品质分组(凡/良/精标题)实锤；当前quality=ring_group故数量=10+全部",
    "TC045": "部分实现：17:10版已改TogContent ToggleContainer+VList虚拟滚动；预览chunk仍旧版(TypeList)未编译生效(D1)",
    "TC046": "部分实现：17:10版已补交互——点击最高档未达标→showMessageTip(『炼宝等级达到7级可查阅』，atlas_red_god_show_cond=7)+勾选回弹；但条件仅maxQuality(仙品)一档，非每ring_group独立；预览旧chunk未取证",
    "TC047": "部分实现：当前按quality筛选，未按ring_group；筛选图已数据化atlas_filter_res(L134引用)",
    "TC048": "待运行时：切换/取消/滚动边界，受预览混杂阻塞(D1)",
    "TC049": "资源就绪：fabao_6020_1/fabao_6045_1/path_to_strength_img04_title19-20/lihui11-12均存在；引用待运行时",
    "TC050": "资源就绪：wukong_gift_img01_fabao01/02_title01均存在；调色落实待运行时",
    "TC051": "资源就绪：completiong_gift_pack_img02_name01~04存在；文案引用待运行时",
    "TC052": "资源就绪：hunhuan图8张+img02_01~03存在；cfg_charge_push_gift_pack无『灵力』；prefab文案待运行时",
    "TC053": "资源就绪；文案待运行时",
    "TC054": "PASS(运行时)：法宝主界面UISoulRingDraw打开正常(v1 17:18/v2 16:07双轮通过)",
    "TC055": "FAIL(运行时)：炼宝页签UISoulRingDrawTabView两轮均未打开(工具上下文/时机限制，非必然产品缺陷，需人工复核)",
    "TC056": "PASS(运行时)：图鉴打开正常(v1/v2双轮通过；probe确认节点UIRootCamera/System/UISoulRingAtlas)",
    "TC057": "PASS(运行时)：法宝装配界面UIHeroSoulRing打开正常",
    "TC058": "待运行时：装配/卸下/替换/一键等操作链，本轮未覆盖",
    "TC059": "待运行时：技能tips打开，本轮未单独覆盖",
    "TC060": "待运行时：前往获取跳转(jump_id)，本轮未单独覆盖；运行时日志见jump_id=40/57/58等配置读取失败告警，建议复核",
    "TC061": "待运行时：前往装配跳转，本轮未单独覆盖",
    "TC062": "PASS(运行时)：概率弹窗UISoulRingDrawLvDesc打开正常",
    "TC063": "FAIL(运行时)：GM面板(节点名uiGM)两轮均未命中(工具限制)",
    "TC064": "待运行时：英雄/伙伴详情法宝页签，本轮未单独覆盖",
    "TC065": "待运行时：战斗内法宝技能释放，本轮未覆盖",
    "TC066": "静态关注+运行时佐证：代码大量直接索引quality(UISoulRing L51等)；运行时日志实际出现『cfg_soul_ring_quality 索引为：-1』『cfg_soul_ring_pool 索引为：0』报错，建议补防御",
    "TC067": "待运行时：ring_group越界处理，受预览混杂阻塞(D1)",
    "TC068": "PASS(静态)：组6(神魂灵/彩)特殊逻辑保留(formatYear/formatLevel对ring_group==6分支仍在)",
    "TC069": "待运行时：版本更新前后数据无缝，受预览混杂阻塞(D1)",
    "TC070": "待运行时：卡池等级边界(pool_level=1实测)，受预览混杂阻塞(D1)",
    "TC071": "待运行时：图鉴滚动帧率；17:10已引入VList虚拟滚动，需新预览验证",
    "TC072": "待运行时：弱网/断网，受预览混杂阻塞(D1)",
    "TC073": "待运行时：多分辨率，受预览混杂阻塞(D1)",
}

STATUS = {
    "PASS": ("通过", "FFC6EFCE"), "部分实现": ("部分实现", "FFFFEB9C"),
    "未实现": ("未实现", "FFFFC7CE"), "无法验证": ("无法验证(需造数)", "FFD9D9D9"),
    "待运行时": ("待运行时(阻塞)", "FFD9D9D9"), "资源就绪": ("资源就绪", "FFFFEB9C"),
    "FAIL": ("失败", "FFFFC7CE"), "静态关注": ("静态关注", "FFFFEB9C"),
}

def status_of(conc: str):
    for k, v in STATUS.items():
        if conc.startswith(k):
            return v
    return ("其他", "FFFFFFFF")

wb = openpyxl.load_workbook(SRC)
ws = wb["测试用例"]

hdr_row = 1
id_col = 1
for c in range(1, ws.max_column + 1):
    if ws.cell(row=hdr_row, column=c).value == "用例编号":
        id_col = c
        break

base_col = ws.max_column
for off, (title, width) in enumerate([("复测结论(trae·全量重测 2026-09-08)", 80), ("结果", 14)], start=1):
    col = base_col + off
    cell = ws.cell(row=hdr_row, column=col, value=title)
    cell.font = Font(bold=True)
    cell.fill = PatternFill("solid", fgColor="FFE9C46A")
    cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = width

n_filled = 0
for row in range(hdr_row + 1, ws.max_row + 1):
    cid = ws.cell(row=row, column=id_col).value
    if not cid:
        continue
    cid = str(cid).strip()
    conc = CONC.get(cid, "")
    if not conc:
        continue
    n_filled += 1
    c1 = ws.cell(row=row, column=base_col + 1, value=conc)
    c1.alignment = Alignment(wrap_text=True, vertical="top")
    label, color = status_of(conc)
    c2 = ws.cell(row=row, column=base_col + 2, value=label)
    c2.fill = PatternFill("solid", fgColor=color)
    c2.alignment = Alignment(vertical="center", horizontal="center")

wb.save(DST)
print("已生成:", DST)
print("回填用例数:", n_filled)

from collections import Counter
stat = Counter()
for v in CONC.values():
    stat[status_of(v)[0]] += 1
total = sum(stat.values())
print("结论分布:", dict(stat))
print("总计:", total)
for k in ["通过", "部分实现", "未实现", "失败", "无法验证(需造数)", "待运行时(阻塞)", "资源就绪", "静态关注"]:
    if k in stat:
        print(f"  {k}: {stat[k]} ({stat[k]/total*100:.0f}%)")
