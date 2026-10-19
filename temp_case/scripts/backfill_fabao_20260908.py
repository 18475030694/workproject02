# -*- coding: utf-8 -*-
"""法宝UI优化 测试用例回填（2026-09-08 复测）——在『测试用例』sheet 追加复测结论列"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

SRC = r"E:\workproject\temp_case\cases\法宝UI优化测试用例.xlsx"
DST = r"E:\workproject\temp_case\cases\法宝UI优化测试用例_已回填.xlsx"

# 每条用例的复测结论（静态权威 + 运行时佐证）
CONC = {
    "TC001": "部分实现：资源已数据化到cfg_soul_ring_quality，但代码仍按quality取道具框(UISoulRing.ts L51-52/83-84)；配置quality==ring_group 100%，显示巧合等价",
    "TC002": "部分实现：光球draw_ball_eff_res已配10档，但代码按quality取；同上巧合等价",
    "TC003": "部分实现：漩涡eff_res已配10档，但代码按quality取(frame_eff_res来自qualityCfg)",
    "TC004": "部分实现：技能框skill_res已配10档，但取用仍按quality",
    "TC005": "部分实现：图鉴框atlas_res已配10档，但取用仍按quality",
    "TC006": "未实现：名称色值仍按quality(cfg_soul_ring_quality[quality].color)，未按ring_group；需求名称色值本身标『等美术给』",
    "TC007": "无法验证：配置quality==ring_group 100%，需GM造quality≠ring_group数据才能暴露",
    "TC008": "无法验证：同上，当前数据无法区分",
    "TC009": "无法验证：代码读quality，改quality必然影响外观",
    "TC010": "PASS(静态)：10组资源文件齐全(bg/光球/漩涡/技能框/图鉴框/筛选图均存在)",
    "TC011": "待运行时：cfg同时保留quality与ring_group字段，数据兼容；运行时受预览不一致阻塞",
    "TC012": "PASS(静态)：UISoulRing列表项已用formatLevel输出阶(Year_txt)",
    "TC013": "PASS(静态)：UISoulRingShow已用formatLevel",
    "TC014": "PASS(静态)：UISoulRingAtlasDesc L52已改『初始阶数』+formatLevel；运行时预览旧缓存仍显『1000年』",
    "TC015": "PASS(静态)：UIRingUnlockSkill L61用formatLevel输出阶",
    "TC016": "PASS(静态)：UIRingUnlockFeatures L59用formatLevel",
    "TC017": "PASS(静态)：UISoulRingDrawRewards/UIRingTipsTemp已用formatLevel",
    "TC018": "PASS(静态)：抽奖结果列表已改阶(UIHeroSoulItem/UIHeroSoulRing L526/563)",
    "TC019": "PASS(静态)：UISoulRingUpgradeResult L112用formatLevel",
    "TC020": "PASS(静态)：RingAtlasItem L85用formatLevel",
    "TC021": "PASS(静态)：推荐条目同列表项链路，已改阶",
    "TC022": "PASS(静态)：UIHeroSoulItem L90/169、UIHeroSoulRing L526已改阶",
    "TC023": "部分实现：SoulRingCtl L69仍用formatYear2(图标档位索引，非文案)，战斗文案待运行时",
    "TC024": "PASS(静态+运行时)：soulring业务代码已无灵力文案(formatYear成死代码)，运行时未见『灵力/万灵力』",
    "TC025": "未实现：阶数颜色仍按quality(noFeatures=true)或灰#8999a1，未按level取色",
    "TC026": "PASS(静态)：数据层已按ring_group定位升级表(SoulRingMgr L369/372)，阶数与加成计算一致",
    "TC027": "PASS(静态)：解锁条件文案已随formatLevel改阶(UIRingUnlockSkill/UIRingUnlockFeatures)",
    "TC028": "待运行时：达到阶数即时解锁依赖服务端事件，运行时受预览不一致阻塞",
    "TC029": "未实现：代码无『未解锁』/『暂无特性』文案；UIHeroSoulItem显示实际效果文案(灰色)而非『未解锁』",
    "TC030": "未实现：无『暂无特性』文案与判定逻辑",
    "TC031": "未实现：特性显示逻辑未按需求实现，边界无法正确验证",
    "TC032": "未实现：同上",
    "TC033": "未实现：各界面特性显示规则不一致(UISoulRingDesc显示全部/UIHeroSoulItem显示全部/UIRingUnlockFeatures只显示已激活)",
    "TC034": "未实现：UISoulRingDesc L123-151显示全部特性(灰/亮区分)，未『只显示已激活』",
    "TC035": "未实现：无『无已激活特性』特殊表现逻辑",
    "TC036": "待运行时：全部激活时显示完整，运行时受预览不一致阻塞",
    "TC037": "未实现：cfg_soul_ring_pool_lv_desc仍year(品质名)+color色值；UISoulRingDrawLvDesc L121/146仍用e.year+e.color",
    "TC038": "未实现：未新增组别底图资源字段",
    "TC039": "未实现：色值配置未废弃仍生效",
    "TC040": "PASS(静态)：升级预览LV.1→LV.2概率数值链路正常(cfg_soul_ring_pool_lv_desc)；运行时弹窗可打开",
    "TC041": "部分实现：法宝列表项UISoulRing无攻防血；装配item UIHeroSoulItem L99-110有基础属性显示",
    "TC042": "待运行时：属性与角色面板一致性需运行时比对，受预览不一致阻塞",
    "TC043": "待运行时：名称上移/长文本遮挡需运行时，受预览不一致阻塞",
    "TC044": "部分实现：分类仍按quality(qualityList来自cfg_soul_ring_quality)，未按ring_group；当前quality=ring_group故数量=10",
    "TC045": "部分实现：TypeList为Toggle容器，滚动待运行时验证",
    "TC046": "部分实现：卡池条件仅maxQuality(仙品)一档(atlas_red_god_show_cond)，非每ring_group独立；运行时未捕获提示",
    "TC047": "部分实现：当前按quality筛选，未按ring_group",
    "TC048": "待运行时：切换/取消/滚动边界，受预览不一致阻塞",
    "TC049": "资源就绪：fabao_6020_1/fabao_6045_1/path_to_strength_img04_title19-20/lihui11-12均存在；引用待运行时",
    "TC050": "资源就绪：wukong_gift_img01_fabao01/02_title01均存在；调色落实待运行时",
    "TC051": "资源就绪：completiong_gift_pack_img02_name01~04存在；文案引用待运行时",
    "TC052": "资源就绪：hunhuan图8张+img02_01~03存在；cfg_charge_push_gift_pack无『灵力』；prefab文案待运行时",
    "TC053": "资源就绪；文案待运行时",
    "TC054": "PASS(运行时)：法宝主界面UISoulRingDraw打开正常",
    "TC055": "FAIL(运行时)：炼宝页签UISoulRingDrawTabView本轮未打开(需上下文/时机调整)",
    "TC056": "PASS(运行时)：图鉴打开正常",
    "TC057": "PASS(运行时)：法宝装配界面UIHeroSoulRing打开正常",
    "TC058": "待运行时：装配/卸下/替换/一键等操作链，本轮未覆盖",
    "TC059": "待运行时：技能tips打开，本轮未单独覆盖",
    "TC060": "待运行时：前往获取跳转(jump_id)，本轮未单独覆盖",
    "TC061": "待运行时：前往装配跳转，本轮未单独覆盖",
    "TC062": "PASS(运行时)：概率弹窗UISoulRingDrawLvDesc打开正常",
    "TC063": "FAIL(运行时)：GM面板(节点名uiGM)本轮未命中",
    "TC064": "待运行时：英雄/伙伴详情法宝页签，本轮未单独覆盖",
    "TC065": "待运行时：战斗内法宝技能释放，本轮未覆盖",
    "TC066": "静态关注：代码大量直接索引quality(UISoulRing L51等)，缺字段会异常；建议补防御",
    "TC067": "待运行时：ring_group越界处理，受预览不一致阻塞",
    "TC068": "PASS(静态)：组6(神魂灵/彩)特殊逻辑保留(formatYear/formatLevel对ring_group==6分支仍在)",
    "TC069": "待运行时：版本更新前后数据无缝，受预览不一致阻塞",
    "TC070": "待运行时：卡池等级边界，受预览不一致阻塞",
    "TC071": "待运行时：图鉴滚动帧率，受预览不一致阻塞",
    "TC072": "待运行时：弱网/断网，受预览不一致阻塞",
    "TC073": "待运行时：多分辨率，受预览不一致阻塞",
}

wb = openpyxl.load_workbook(SRC)
ws = wb["测试用例"]

# 找到表头行，定位用例编号列与最后一列
hdr_row = 1
for r in range(1, 3):
    for c in range(1, ws.max_column + 1):
        v = ws.cell(row=r, column=c).value
        if v == "用例编号":
            hdr_row = r
            break

id_col = None
for c in range(1, ws.max_column + 1):
    if ws.cell(row=hdr_row, column=c).value == "用例编号":
        id_col = c
        break

# 追加一列：复测结论(2026-09-08)
new_col = ws.max_column + 1
title = "复测结论(2026-09-08)"
ws.cell(row=hdr_row, column=new_col, value=title)
ws.cell(row=hdr_row, column=new_col).font = Font(bold=True)
ws.cell(row=hdr_row, column=new_col).fill = PatternFill("solid", fgColor="FFE9C46A")
ws.cell(row=hdr_row, column=new_col).alignment = Alignment(wrap_text=True, vertical="center")

fill_map = {
    "PASS": "FFC6EFCE", "部分": "FFFFEB9C", "未实现": "FFFFC7CE",
    "无法验证": "FFD9D9D9", "待运行": "FFD9D9D9", "资源就绪": "FFFFEB9C",
    "FAIL": "FFFFC7CE", "静态关注": "FFFFEB9C",
}

for row in range(hdr_row + 1, ws.max_row + 1):
    cid = ws.cell(row=row, column=id_col).value
    if not cid:
        continue
    cid = str(cid).strip()
    concl = CONC.get(cid, "")
    cell = ws.cell(row=row, column=new_col, value=concl)
    cell.alignment = Alignment(wrap_text=True, vertical="top")
    # 按结论着色
    for k, color in fill_map.items():
        if concl.startswith(k):
            cell.fill = PatternFill("solid", fgColor=color)
            break

ws.column_dimensions[openpyxl.utils.get_column_letter(new_col)].width = 70

wb.save(DST)
print("已回填:", DST)
print("回填用例数:", len(CONC))

# 统计
from collections import Counter
stat = Counter()
for c, v in CONC.items():
    if v.startswith("PASS"): stat["通过"] += 1
    elif v.startswith("部分"): stat["部分实现"] += 1
    elif v.startswith("未实现"): stat["未实现"] += 1
    elif v.startswith("无法验证"): stat["无法验证(需造数)"] += 1
    elif v.startswith("待运行"): stat["待运行时(阻塞)"] += 1
    elif v.startswith("资源就绪"): stat["资源就绪"] += 1
    elif v.startswith("FAIL"): stat["失败"] += 1
    else: stat["其他"] += 1
print("结论分布:", dict(stat))
