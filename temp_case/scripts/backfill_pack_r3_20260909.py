# -*- coding: utf-8 -*-
"""法宝用例回填（2026-09-09 公告滚屏R3 + 推送礼包取证）——更新 v2 Excel 的 TC050~053。

依据：
  - reports/fabao_pack_ui_evidence_20260909_192734/（5 UI 打开 + 截图 + 文本 dump）
  - reports/fabao_r3_bulletin_20260909_191601/（R3 公告 9/9 PASS，73 条用例无公告映射，仅入报告）
  - 源码：UIRechargeSoulRing.ts L121 res1 引用 / cfg_soulring_push_gift 480001~4
"""
import openpyxl

XLSX = r"E:\workproject\temp_case\cases\法宝UI优化测试用例_trae_v2.xlsx"

# TC -> (复测结论列更新, 结果列更新)
UPD = {
    "TC050": (
        "PASS(运行时 09-09)：UIWuKongGift 打开正常，紫金葫芦/筋斗云标签实显『极品』（原『玄品』已改）；"
        "调色/光晕细节为图片资源，建议人工对比需求图。证据 reports/fabao_pack_ui_evidence_20260909_192734/04_wukong_gift.png",
        "通过(文案实锤)",
    ),
    "TC051": (
        "PASS(运行时 09-09)：UIRechargeSoulRing 标题图实显『150阶法宝 [无始钟]』新图；"
        "res1=completiong_gift_pack_img02_name01~04 经 UIRechargeSoulRing.ts L121 loadSetUIIcon 引用生效；"
        "4 档[圣]无始钟/昆吾人皇剑/番天印/创世鼎页签齐全、￥98。证据 …/05_recharge_soulring.png",
        "通过(截图实锤)",
    ),
    "TC052": (
        "PASS(运行时 09-09)：UIPushGiftPack/UIRechargeSoulRing/UIWuKongGift/UIGodStarGift 全量文本 dump "
        "0『灵力』命中（含推送礼包浮层）；hunhuan 8 张阶数图为图片内文案，逐张对比建议人工复核",
        "通过(0残留)",
    ),
    "TC053": (
        "待人工复核：UIGodStarGift 打开正常，但 24 档仙星礼包已于前置轮次全部购买"
        "（recharge_notifykey 24/24 OK），内容区为空，img02_01~03 文案图需新号/未购账号复核",
        "待人工复核(已购空)",
    ),
}

wb = openpyxl.load_workbook(XLSX)
ws = wb["测试用例"]

hdr_row = 1
id_col = concl_col = res_col = None
for c in range(1, ws.max_column + 1):
    v = str(ws.cell(row=hdr_row, column=c).value or "")
    if v == "用例编号":
        id_col = c
    elif v.startswith("复测结论"):
        concl_col = c
    elif v == "结果":
        res_col = c
assert id_col and concl_col and res_col, (id_col, concl_col, res_col)

hit = 0
for r in range(2, ws.max_row + 1):
    tid = str(ws.cell(row=r, column=id_col).value or "").strip()
    if tid in UPD:
        concl, res = UPD[tid]
        ws.cell(row=r, column=concl_col, value=concl)
        ws.cell(row=r, column=res_col, value=res)
        hit += 1
        print(f"[ok] {tid} -> {res}")

wb.save(XLSX)
print(f"saved: {XLSX} (updated {hit}/4)")
