"""大圣礼包 E2E 结果回填：result_*.json → 用例表（天数锚点版）+ Markdown 报告。

用法：python temp_case/scripts/backfill_wukong_results.py
"""
import json
import re
import sys
import time
from datetime import datetime
from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve().parent
REPORTS = HERE.parent / "reports"
CASES = HERE.parent / "cases"
XLSX = CASES / "大圣礼包调整测试用例_天数锚点版.xlsx"
MD = REPORTS / "wukong_gift_e2e_report.md"

GROUP_TITLES = {
    "ga": "A组·新号购买流", "gb": "B组·D+1跨日", "gc": "C组·D+2跨日",
    "a2": "A2组·漏领/全领", "e01": "E组·0点边界", "compat": "E组·已购兼容",
    "expiry": "F组·过期占位", "state": "STATE", "claim": "CLAIM",
    "rejudge": "复判·伪失败纠正（含配置核对+魂骨实测）",
}


def norm_tc(key):
    """'TC-016a#2402(重复购买拦截)' → 'TC-016'"""
    m = re.match(r"(TC-[A-Z]?\d+)", key)
    return m.group(1) if m else None


def load_results():
    subs = {}  # (tc, sub) -> rec（按目录时间序，后跑覆盖先跑：同子项取最新）
    runs = []  # group-level meta
    for d in sorted(REPORTS.glob("wukong_gift_e2e_2*")):
        for f in sorted(d.glob("result_*.json")):
            try:
                data = json.loads(f.read_text(encoding="utf-8"))
            except Exception:  # noqa: BLE001
                continue
            g = data.get("group", f.stem[7:])
            runs.append({"dir": d.name, "group": g, "time": data.get("time"),
                         "server_time": data.get("server_time"),
                         "extra": {k: v for k, v in data.items()
                                   if k in ("account", "uuid")}})
            for r in data.get("results", []):
                tc = norm_tc(r.get("tc", ""))
                if not tc:
                    continue
                subs[(tc, r.get("tc"))] = {
                    "sub": r.get("tc"), "ok": r.get("pass"), "note": r.get("note", ""),
                    "evidence": r.get("evidence") or [], "group": g,
                    "dir": d.name, "ts": data.get("time", ""),
                }
    agg = {}
    for (tc, _sub), rec in subs.items():
        agg.setdefault(tc, []).append(rec)
    return agg, runs


def fill_xlsx(agg):
    wb = openpyxl.load_workbook(XLSX)
    filled = []
    table_tcs = []
    for sheet in ("测试用例", "异常与边界用例"):
        if sheet not in wb.sheetnames:
            continue
        ws = wb[sheet]
        for row in ws.iter_rows(min_row=2):
            tc = row[0].value
            if not tc:
                continue
            table_tcs.append(tc)
            if tc not in agg:
                continue
            recs = agg[tc]
            lines = []
            for r in recs:
                mark = "✓" if r["ok"] else "✗"
                ev = f"（证据: {'、'.join(r['evidence'])}）" if r["evidence"] else ""
                lines.append(f"{mark} {r['sub']}: {r['note']}{ev}")
            actual = "\n".join(lines)
            verdict = "PASS" if all(r["ok"] for r in recs) else "FAIL"
            # 截图/证据汇总到备注
            all_ev = [e for r in recs for e in r["evidence"]]
            note = f"{'、'.join(dict.fromkeys(all_ev))}" if all_ev else ""
            row[8].value = actual[:32000]
            row[9].value = verdict
            if note:
                row[10].value = ((row[10].value + "\n") if row[10].value else "") + f"[回填证据] {note}"
            filled.append((tc, verdict, len(recs)))
    # 保存（锁文件保护）
    uncovered = [t for t in table_tcs if t not in agg]
    for attempt in range(13):
        try:
            wb.save(XLSX)
            return filled, str(XLSX), uncovered
        except PermissionError:
            if attempt == 0:
                print(f"[warn] {XLSX.name} 被 Excel 锁定，请关闭后等待自动重试（最长2分钟）…")
            time.sleep(10)
    alt = XLSX.with_name(XLSX.stem + "_回填.xlsx")
    wb.save(alt)
    print(f"[warn] 原表仍锁定，已另存 {alt.name}")
    return filled, str(alt), [t for t in table_tcs if t not in agg]


def write_md(agg, runs, filled, xlsx_used, uncovered):
    now = datetime.now().strftime("%Y-%m-%d %H:%M")
    lines = [f"# 大圣礼包（530001）改版 E2E 测试报告", "",
             f"- 生成时间：{now}",
             f"- 用例表：{xlsx_used}（回填 {len(filled)} 条）",
             f"- 服务器：S5（192.168.*.* GM），部署版 fb_gift（购买锚点 day_no=Date.TodayNo()）",
             f"- 服务器时间：测试中通过 set_timeshift（绝对模型）+ player_daily_reset 跨日（窗口 2026-10-23 ~ 2026-10-29）；测试结束已 restore 拨回",
             f"- 结果目录：temp_case/reports/wukong_gift_e2e_*/", ""]
    lines += ["## 分组执行情况", "", "| 分组 | 时间 | 服务器时间 | 账号 |", "|---|---|---|---|"]
    for r in runs:
        acc = r["extra"].get("account", "-")
        lines.append(f"| {GROUP_TITLES.get(r['group'], r['group'])} | {r['time']} | {r['server_time']} | {acc} |")
    lines += ["", "## 用例结果", "", "| 用例编号 | 结果 | 记录数 | 实际结果摘要 |", "|---|---|---|---|"]
    for tc in sorted(agg):
        recs = agg[tc]
        verdict = "PASS" if all(r["ok"] for r in recs) else "FAIL"
        first = recs[0]["note"][:80].replace("|", "／").replace("\n", " ")
        lines.append(f"| {tc} | {verdict} | {len(recs)} | {first} |")
    lines += ["", "## 发现的问题与风险", ""]
    fails = [tc for tc in agg if not all(r["ok"] for r in agg[tc])]
    if fails:
        for tc in fails:
            for r in agg[tc]:
                if not r["ok"]:
                    lines.append(f"- **{r['sub']}**（{GROUP_TITLES.get(r['group'], r['group'])}）：{r['note']}")
    else:
        lines.append("- 无自动化用例失败。")
    lines += ["", "### 代码审读发现（含实测佐证）", "",
              "- **P0 资损（TC-E05 实测确认）**：过期未购账号（如2402）礼包过期后占位数据（count=0）经 SceneReady 照常下发，客户端出现领取入口；C2gsGetFbGiftReward 无 count 拦截，直发领取成功（day1 置为已领并发放 reward_list1）。建议：SceneReady 过滤 count<=0 + 领取接口拦截。",
              "- **P2 风险**：OnDailyReset 无 giftCfg-nil / day_no-nil 防御，cfg 删除或脏数据时每日重置会报错（当前配置存在未触发）。",
              "- **口径差异**：老存档兼容实现为「首登日锚点，新增奖励按后续自然日逐日解锁（D2=次日0点，D3=再次日）」，与用例备注中『首个0点全解锁未领份额』的兼容口径不一致，按部署版实现验收（TC-022 已按此口径执行）。",
              "- **拨时机制**：本服 set_server_time 的 timestr 解析不可用（编码问题）；set_timeshift?dt= 为**绝对设置**（TIME_SHIFT_SEC := N，非增量），服务器时间 = 真实时间 + 614867 + dt（实测 dt=236284 精确命中目标）。跨0点 = 拨至前一日 23:57:30 由自然时钟跨日 + player_daily_reset；拨时后客户端必须重登才刷新时间。", "",
              "### 未自动化 / 待人工执行用例", "",
              "以下用例不在本次 E2E 自动化范围（需人工交互/多系统联动/主观判断），已在用例表中留空待人工回填：", ""]
    for tc in uncovered:
        lines.append(f"- {tc}")
    MD.write_text("\n".join(lines), encoding="utf-8")
    print(f"[out] {MD}")


def main():
    agg, runs = load_results()
    if not agg:
        print("[err] 未找到任何 result_*.json")
        sys.exit(1)
    filled, xlsx_used, uncovered = fill_xlsx(agg)
    print(f"[ok] 回填 {len(filled)} 条 -> {xlsx_used}（未覆盖 {len(uncovered)} 条）")
    for tc, verdict, n in filled:
        print(f"  {tc}: {verdict} ({n})")
    write_md(agg, runs, filled, xlsx_used, uncovered)


if __name__ == "__main__":
    main()
