# -*- coding: utf-8 -*-
"""新秘宝 A组动画专项 回填：data.json → 覆盖写入 已回填用例表。
规则(用户 2026-09-12): 覆盖前一用例表（不另存新文件）；表格被 Excel 打开占用 → 提醒；
2分钟未响应再提醒并另存新表。用法: python backfill_mibao_ani.py [data_json路径]
退出码: 0=成功 2=文件被占用（需重试或另存）"""
import json
import sys
import time
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

ROOT = Path(r"E:\workproject")
FILLED_XLSX = ROOT / "temp_case" / "cases" / "新秘宝用例_v2执行版_已回填.xlsx"
SOURCE_XLSX = ROOT / "temp_case" / "cases" / "新秘宝用例_v2执行版.xlsx"
REPORT = ROOT / "temp_case" / "reports" / "新秘宝A组动画测试报告.md"
SHEET_ANI = "A组动画专项"

# A组专项用例 → 主表已有行(列B 用例编号) 映射；未列出的只写专项sheet
MAP_TO_MAIN = {
    "SPD-001S": "MIB-G1",
    "ANI-008": "MIB-G5",
    "MIB-G5": "MIB-G5",   # 与 ANI-008 同源，后者优先
    "ABN-008S": "MIB-L7",
    "ABN-010": "MIB-L9",
    "MIB-I8": "MIB-I8",
}

data_path = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "reports" / "e2e_mibao_ani_latest" / "data.json"
data = json.loads(data_path.read_text(encoding="utf-8"))
res = {}
for c in data.get("cases", []):
    res[c["case_id"]] = c  # 同 id 后写覆盖前写

fills = {"PASS": PatternFill("solid", fgColor="C6EFCE"),
         "FAIL": PatternFill("solid", fgColor="FFC7CE"),
         "NA": PatternFill("solid", fgColor="D9D9D9")}


def tag_of(e):
    return "PASS" if e["passed"] is True else ("FAIL" if e["passed"] is False else "NA")


def save_or_new(wb, out_path: Path):
    """覆盖保存；被占用返回 (False, None)"""
    try:
        wb.save(out_path)
        return True, out_path
    except PermissionError:
        return False, None


def save_with_rule(wb, out_path: Path):
    """用户规则(2026-09-12): 占用→提醒→等2分钟重试→仍占用→再提醒并另存带后缀新表"""
    ok, saved = save_or_new(wb, out_path)
    if ok:
        return ok, saved
    print(f"OCCUPIED(表格被Excel打开): {out_path}")
    print("REMIND: 请关闭该 Excel 文件，2分钟后自动重试...")
    time.sleep(120)
    ok, saved = save_or_new(wb, out_path)
    if ok:
        return ok, saved
    print("REMIND: 仍未关闭，另存新表...")
    alt = out_path.with_name(out_path.stem + f"_另存_{datetime.now().strftime('%H%M%S')}.xlsx")
    ok, saved = save_or_new(wb, alt)
    return ok, saved


wb = openpyxl.load_workbook(FILLED_XLSX) if FILLED_XLSX.exists() else openpyxl.load_workbook(SOURCE_XLSX)
ws = wb["测试用例清单"]
HDR = 10  # 模块..实现依据 10列，回填列 11-14
COLS = ["执行结果", "复判", "执行说明", "截图"]

stat = {"PASS": 0, "FAIL": 0, "NA": 0}
for e in res.values():
    if e["case_id"] in ("LOGIN", "OPENUI", "PREP", "EXCEPTION"):
        continue
    stat[tag_of(e)] += 1

# 1) 主表映射行更新
for src_id, main_id in MAP_TO_MAIN.items():
    e = res.get(src_id)
    if not e or src_id == "MIB-G5":  # G5 用 ANI-008 的记录
        continue
    row = None
    for i in range(2, ws.max_row + 1):
        if ws.cell(row=i, column=2).value == main_id:
            row = i
            break
    if not row:
        continue
    t = tag_of(e)
    ws.cell(row=row, column=HDR + 1, value=t).fill = fills[t]
    ws.cell(row=row, column=HDR + 1).alignment = Alignment(horizontal="center")
    ws.cell(row=row, column=HDR + 2, value="待复判" if t == "FAIL" else ("—" if t == "NA" else "通过"))
    ws.cell(row=row, column=HDR + 3, value=f"[A组动画专项] {str(e.get('note') or '')[:400]}")
    ws.cell(row=row, column=HDR + 4, value=str(e.get("screenshot") or ""))

# 2) 专项 sheet（重建）
if SHEET_ANI in wb.sheetnames:
    del wb[SHEET_ANI]
ws2 = wb.create_sheet(SHEET_ANI)
ws2.append(["新秘宝(仙机) A组动画专项执行结果"])
ws2["A1"].font = Font(bold=True, size=13)
ws2.append(["执行时间", data.get("start_time"), "账号", f"{data.get('account')}@{data.get('server')}",
            "产物", str(data.get("out_dir"))])
ws2.append([])
ws2.append(["用例编号", "标题", "执行结果", "执行说明", "截图"])
for c in ws2[4]:
    c.font = Font(bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor="548235")
ANI_ORDER = ["RES-001", "RES-005", "ANI-001", "ANI-002", "ANI-003", "ANI-004", "ANI-005",
             "ANI-006", "ANI-007", "ANI-008", "ANI-009", "ANI-010", "ANI-011",
             "GAM-002", "GAM-004S", "SPD-001S", "ABN-008S", "ABN-010", "MIB-I8"]
for cid in ANI_ORDER:
    e = res.get(cid)
    if not e:
        ws2.append([cid, "", "NA", "本轮未执行", ""])
        continue
    t = tag_of(e)
    ws2.append([cid, e.get("title"), t, str(e.get("note") or "")[:500], str(e.get("screenshot") or "")])
    ws2.cell(row=ws2.max_row, column=3).fill = fills[t]
    ws2.cell(row=ws2.max_row, column=3).alignment = Alignment(horizontal="center")
for col, w in zip("ABCDE", [12, 30, 10, 60, 26]):
    ws2.column_dimensions[col].width = w

# 3) 概要追加
if "E2E概要" in wb.sheetnames:
    wss = wb["E2E概要"]
    wss.append([])
    wss.append([f"A组动画专项({data.get('start_time')}) PASS={stat['PASS']} FAIL={stat['FAIL']} NA={stat['NA']}"])
    wss.cell(row=wss.max_row, column=1).font = Font(bold=True)
    for c in data.get("cases", []):
        if c["case_id"] in ("LOGIN", "OPENUI", "PREP", "EXCEPTION"):
            continue
        wss.append([c["case_id"], c["title"], "FAIL" if c["passed"] is False else ("PASS" if c["passed"] is True else "NA"),
                    str(c.get("note") or "")[:180]])

ok, saved = save_with_rule(wb, FILLED_XLSX)
if not ok:
    print("SAVE-FAILED:", FILLED_XLSX)
    sys.exit(2)

# 4) MD 报告
fails = [c for c in data.get("cases", []) if c["passed"] is False]
nas = [c for c in data.get("cases", []) if c["passed"] is None]
passes = [c for c in data.get("cases", []) if c["passed"] is True and c["case_id"] not in ("LOGIN", "OPENUI", "PREP")]
L = ["# 新秘宝(仙机) A组动画专项测试报告", "",
     f"- 时间: {data.get('start_time')}  账号: {data.get('account')}@{data.get('server')}",
     f"- 汇总: **PASS {stat['PASS']} / FAIL {stat['FAIL']} / NA {stat['NA']}**（19条 = 缓测清单18 + MIB-I8）",
     f"- 产物: {data.get('out_dir')}",
     f"- 回填: 覆盖写入 {FILLED_XLSX.name}（主表映射 G1/G5/L7/L9/I8 + 专项sheet）", "",
     "## FAIL/待复核明细", "", "| 用例 | 标题 | 说明 |", "|---|---|---|"]
L += [f"| {c['case_id']} | {c['title']} | {str(c.get('note') or '')[:200]} |" for c in fails] if fails else ["| — | 无自动判 FAIL | |"]
L += ["", "## 全量结果", "", "| 用例 | 标题 | 结果 | 说明 |", "|---|---|---|---|"]
for cid in ANI_ORDER:
    e = res.get(cid)
    if not e:
        continue
    L.append(f"| {cid} | {e['title']} | {tag_of(e)} | {str(e.get('note') or '')[:160]} |")
REPORT.write_text("\n".join(L), encoding="utf-8")
print("backfill(overwrite):", saved)
print("report :", REPORT)
print("summary:", stat)
