# -*- coding: utf-8 -*-
"""生成《月卡、终身卡调整》测试用例 Excel"""
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

OUT = r"E:\workproject\temp_case\cases\月卡终身卡调整测试用例.xlsx"

# ============ Sheet1 主流程用例 ============
HDR = ["用例编号", "所属模块", "用例标题", "优先级", "前置条件", "测试步骤", "预期结果", "关联需求", "实际结果", "测试结果", "备注"]

cases = [
    # 模块一：处理1 弹窗替换
    ["TC-001", "模块一：弹窗替换（处理1）", "月卡解锁时挂机界面主动弹出月卡购买界面", "P0",
     "1.新号未购买月卡（card 2 buy_time=0）\n2.主线进度<40关\n3.停留在挂机主界面",
     "①推进主线通关第40关，触发月卡系统(sys 87)解锁\n②观察弹窗内容",
     "1.挂机界面主动弹出月卡购买解锁界面UIPrivCard（cardId=2）\n2.样式与需求图1一致：标题'月卡'、'30天特权'、购买即刻得、每日可领、特权列表（每日飞速推进+1/挂机奖励收集+2小时/主线2倍速+1/仙机副本+1/炼宝副本+1/武器副本+1）、底部¥18购买按钮",
     "处理1", "", "", ""],
    ["TC-002", "模块一：弹窗替换（处理1）", "月卡解锁时不再弹出终身卡宣传界面", "P0",
     "同TC-001",
     "①通关主线第40关解锁月卡/终身卡系统\n②观察全部弹窗序列",
     "1.全程不弹出终身卡宣传界面UIForeverCardOpenAlert\n2.月卡系统解锁仅弹UIPrivCard",
     "处理1", "", "", "代码现状：PrivilegeMgr.ts L150-167终身卡分支已注释"],
    ["TC-003", "模块一：弹窗替换（处理1）", "已购月卡玩家解锁时不再弹窗", "P0",
     "1.月卡在有效期内（buy_time>0且未过期）\n2.主线进度<40关",
     "①通关主线第40关",
     "1.不弹出UIPrivCard\n2.也不弹终身卡宣传界面",
     "处理1", "", "", "getPrivilegeCardStatus(2)=true直接return"],
    ["TC-004", "模块一：弹窗替换（处理1）", "已购终身卡且未购月卡仍弹月卡", "P1",
     "1.已购终身卡（card 3 buy_time>0）\n2.月卡未购买",
     "①通关主线第40关\n②观察弹窗",
     "1.仍主动弹出月卡购买界面UIPrivCard(2)\n2.终身卡宣传弹窗不出现",
     "处理1", "", "", ""],
    ["TC-005", "模块一：弹窗替换（处理1）", "弹窗仅触发一次，重登不重复骚扰", "P1",
     "已触发过一次月卡弹窗并关闭",
     "①关闭月卡弹窗\n②重新登录\n③继续推进主线/停留在挂机界面",
     "1.同一解锁事件弹窗只出现一次\n2.重新登录后不再主动弹出",
     "处理1", "", "", "onNewSysOpen仅在MSG_ON_NEW_SYS_OPEN事件触发"],
    ["TC-006", "模块一：弹窗替换（处理1）", "非主界面时弹窗延迟到回主界面再弹", "P1",
     "月卡即将解锁",
     "①通关第40关时正处于战斗/子界面中\n②返回挂机主界面",
     "1.不在战斗/子界面上强行弹出\n2.回到主界面后再弹出月卡购买界面",
     "处理1", "", "", "队列参数TipShowCondition.OnlyMain"],

    # 模块二：处理1回归
    ["TC-007", "模块二：弹窗回归", "离线闯关特权卡/体验卡弹窗不受影响", "P1",
     "离线闯关特权卡(sys 97)未解锁账号",
     "①触发离线闯关特权卡/体验卡解锁\n②观察弹窗",
     "1.仍正常弹出UIPrivCardOpenAlert\n2.体验卡(sys 142)开服2天内限制逻辑不变",
     "处理1-回归", "", "", "PrivilegeMgr.ts L186-237"],
    ["TC-008", "模块二：弹窗回归", "终身卡详情界面可正常打开", "P1",
     "终身卡已解锁",
     "①通过特权卡入口打开终身卡详情UIPrivCard(3)",
     "1.正常显示终身卡详情（与需求图2一致：'终身卡'、'永久专属特权'、购买即刻得、每日可领、特权列表、¥128按钮）\n2.无异常报错",
     "处理1-回归", "", "", ""],
    ["TC-009", "模块二：弹窗回归", "UIForeverCardOpenAlert无残留触发入口", "P2",
     "代码走查（已确认）",
     "①全局检索UIForeverCardOpenAlert引用\n②实机验证各入口",
     "1.仅剩UIConfig.ts注册（L3145-3148）与PrivilegeMgr注释代码\n2.无任何主动打开入口，玩家无法再触达该界面",
     "处理1-回归", "", "", "静态检索已确认，实机抽查即可"],

    # 模块三：处理2 首购展示
    ["TC-010", "模块三：首购奖励展示（处理2）", "月卡未购时显示首购额外奖励组件", "P0",
     "月卡未购买（buy_time=0）",
     "①打开月卡购买界面UIPrivCard(2)",
     "1.显示首购文案'首次购买额外可得神话仙友随机礼包*1'（来自first_desc字段）\n2.下方显示奖励图标：神话仙友随机礼包×1（20013，来自first_reward字段）\n3.样式与需求图1一致",
     "处理2", "", "", "cfg_privilege_card[2]: first_desc/first_reward"],
    ["TC-011", "模块三：首购奖励展示（处理2）", "终身卡未购时显示首购额外奖励组件", "P0",
     "终身卡未购买（buy_time=0）",
     "①打开终身卡详情UIPrivCard(3)",
     "1.显示首购文案（来自first_desc字段）\n2.下方显示5个奖励图标：攻/敏/防/控/辅·自选仙友碎片各10（20018~20022，共50）\n3.样式与需求图2比对",
     "处理2", "", "",
     "差异：配置文案为'购买额外可得神话仙友随机碎片*50'，需求图为'首次购买额外可得神话仙友随机碎片'（无*50）→见待确认1"],
    ["TC-012", "模块三：首购奖励展示（处理2）", "未配置首购字段的卡不显示额外奖励组件", "P0",
     "免广告卡(1)/乾坤藏甲(4)/药园(5)/灵脉(6)/离线闯关(7)/灵宠(8)/觉醒(9)/藏品(11)任一未购买",
     "①逐个打开上述特权卡详情界面",
     "1.均不显示首购额外奖励组件节点（FirstDescLb/FirstRewardList隐藏）\n2.界面布局无空洞、无错位",
     "处理2", "", "", "这些卡配置中无first_desc/first_reward字段"],
    ["TC-013", "模块三：首购奖励展示（处理2）", "首购后额外奖励组件隐藏", "P0",
     "已首购月卡（或终身卡）",
     "①购买成功后重新打开月卡/终身卡详情界面",
     "1.首购文案与首购奖励图标均隐藏\n2.'购买即刻得'图标加已获得蒙版（setMaskRewared）\n3.其余组件显示正常",
     "处理2", "", "", "UIPrivCard.ts updataFirstData L385/390：info.buy_time==0才显示"],

    # 模块四：处理2 首购发放
    ["TC-014", "模块四：首购奖励发放（处理2）", "月卡首购订单回调发放首购奖励并弹奖励弹窗", "P0",
     "1.月卡未购买\n2.可GM模拟充值（notify_key=PrivilegeCharge.2）",
     "①月卡界面点击¥18购买完成支付（或GM：HostURL/index/recharge_notifykey?uuid=玩家ID&notify_key=PrivilegeCharge.2）\n②观察奖励弹窗\n③核对背包",
     "1.充值订单完成回调时发放（非点击按钮立即发）\n2.奖励弹窗一次性展示：每日奖励(元宝100+炼宝玉佩30，仅首日)+首购奖励(神话仙友随机礼包×1)+购买奖励(元宝180+全服红包×1)\n3.背包对应道具数量正确到账",
     "处理2", "", "", "mod_privilege.lua L380-382首购奖励并入rewardList统一发放"],
    ["TC-015", "模块四：首购奖励发放（处理2）", "终身卡首购订单回调发放首购奖励并弹奖励弹窗", "P0",
     "1.终身卡未购买\n2.GM模拟充值（notify_key=PrivilegeCharge.3）",
     "①完成终身卡购买支付（或GM模拟充值）\n②观察奖励弹窗\n③核对背包",
     "1.奖励弹窗展示：每日奖励(元宝200+凡灵召符30)+首购奖励(5种自选仙友碎片各10)+购买奖励(柔兔梦语×1+元宝1280+全服红包×1)\n2.隐藏奖励buy_reward_hide(10×1,18×1,19×1,16×10,4×10)静默到账、不出现在弹窗\n3.背包数量正确",
     "处理2", "", "", "send=1：购买时另有系统消息+跑马灯"],
    ["TC-016", "模块四：首购奖励发放（处理2）", "月卡有效期内重复购买不再获得首购奖励", "P0",
     "已首购月卡且在有效期内",
     "①再次购买月卡（叠加时长）\n②观察奖励弹窗",
     "1.不再发放首购奖励（无神话仙友随机礼包）\n2.不再发放每日奖励（daily_reward_state已置1）\n3.仅发放购买奖励(元宝180+全服红包×1)\n4.剩余时间+30天，显示'再次购买可延长时间'",
     "处理2", "", "", "isFirstBuy=GetBuyTimes()==0，重复购买时为false"],
    ["TC-017", "模块四：首购奖励发放（处理2）", "月卡过期后续购不发首购奖励", "P0",
     "月卡已过期（历史购买过1次）",
     "①再次购买月卡\n②观察奖励弹窗",
     "1.不发首购奖励\n2.发购买奖励+新周期每日奖励（首日）\n3.时长重新计30天",
     "处理2", "", "", "仅首次购买可获得，与是否过期无关"],
    ["TC-018", "模块四：首购奖励发放（处理2）", "终身卡不可重复购买", "P0",
     "已购买终身卡",
     "①打开终身卡详情界面\n②尝试再次购买（实机/GM回调重放）",
     "1.客户端详情页购买按钮隐藏、显示'已激活'标识\n2.服务端收到重复订单回调返回ALREADY_HAD_PRIVILEGE_CARD，不重复发奖",
     "处理2", "", "", "time_limit=-1已购直接拒绝"],
    ["TC-019", "模块四：首购奖励发放（处理2）", "首购奖励到账数量精确校验", "P1",
     "刚完成首购",
     "①打开背包逐项核对数量",
     "1.月卡：神话仙友随机礼包(20013)×1\n2.终身卡：攻/敏/防/控/辅·自选仙友碎片(20018~20022)各10，共50\n3.与cfg_privilege_card.first_reward完全一致",
     "处理2", "", "", ""],

    # 模块五：回归
    ["TC-020", "模块五：回归测试", "月卡/终身卡每日奖励领取回归", "P1",
     "已购买月卡或终身卡",
     "①次日登录领取每日奖励\n②检查红点",
     "1.每日奖励金额与配置一致（月卡：元宝100+炼宝玉佩30；终身卡：元宝200+凡灵召符30）\n2.领取后红点消失、reward_state同步",
     "回归", "", "", ""],
    ["TC-021", "模块五：回归测试", "特权卡入口打开月卡/终身卡详情UI比对", "P1",
     "特权功能已解锁",
     "①从特权卡入口分别打开月卡/终身卡详情\n②与需求图1/图2逐区域比对",
     "1.标题、卡面背景、'30天特权'/'永久专属特权'、购买即刻得、每日可领、特权列表、首购奖励、价格按钮（¥18/¥128）位置与图文一致\n2.注意：需求图数值(160/1250)与当前配置(180/1280)不一致→见待确认2",
     "回归", "", "", ""],
    ["TC-022", "模块五：回归测试", "月卡剩余时长与续购文案显示", "P1",
     "已购月卡",
     "①打开月卡详情\n②观察剩余时间显示",
     "1.正确显示剩余时间倒计时\n2.显示'剩余X天，再次购买可延长时间'\n3.剩余<1天时倒计时秒级刷新",
     "回归", "", "", "UIPrivCard.ts L261-273"],
    ["TC-023", "模块五：回归测试", "特权卡红点逻辑回归", "P2",
     "已购买且有可领每日奖励",
     "①观察特权卡入口红点\n②领取后再次观察",
     "1.可领时红点亮起（red_point 1093月卡/1094终身卡）\n2.领取后红点消失",
     "回归", "", "", "PrivilegeMgr.updateRedPoint"],
]

# ============ Sheet2 异常与边界用例 ============
exc_cases = [
    ["TC-E01", "异常-界面健壮性", "弱网/未同步时打开月卡界面防崩溃", "P1",
     "可模拟弱网（高延迟/断线重连）",
     "①弱网环境下断线重连瞬间快速打开月卡/终身卡详情界面\n②观察是否报错",
     "1.界面正常打开或安全降级，不崩溃\n2.无'Cannot read properties of undefined'类报错",
     "", "", "",
     "风险点：UIPrivCard.ts L385/390 updataFirstData中info.buy_time无空值保护，若cardInfoMap未同步info为undefined会崩溃，重点验证"],
    ["TC-E02", "异常-交易安全", "购买按钮快速连点防重复下单", "P0",
     "月卡/终身卡未购买",
     "①快速连续点击购买按钮10~20次\n②观察订单与奖励",
     "1.仅生成一次支付订单/一次订单回调\n2.首购奖励仅发放一次，无重复\n3.界面按钮有防连点处理",
     "", "", "", ""],
    ["TC-E03", "异常-交易安全", "订单回调丢失/延迟不发奖不补发错乱", "P1",
     "可模拟支付中断（杀进程/断网）",
     "①点击购买后支付未完成回调即杀进程\n②重新登录检查奖励与卡状态",
     "1.未完成订单不发首购奖励、不激活卡\n2.重新支付成功后仍按首购处理（buy_times未+1）\n3.无刷首购奖励漏洞",
     "", "", "", "服务端以订单回调为准，客户端状态不落盘"],
    ["TC-E04", "异常-边界", "背包已满时首购奖励不丢失", "P2",
     "背包/道具栏已满",
     "①完成月卡首购\n②检查奖励去向",
     "1.奖励转邮箱发放或按背包溢出规则处理\n2.道具不丢失、数量正确",
     "", "", "", ""],
    ["TC-E05", "异常-配置容错", "first_reward配置为空数组时界面容错", "P2",
     "GM/测试环境将某卡first_reward改为空数组[]",
     "①打开该卡详情界面\n②购买该卡",
     "1.首购组件不显示（FirstRewardList.active=false）、无报错\n2.购买流程正常",
     "", "", "", "Table.Empty(cfg.first_reward)服务端跳过；客户端!!cfg.first_reward对空数组为true但循环0次，界面需无空洞"],
    ["TC-E06", "异常-边界", "跨天时点首购发放正确", "P2",
     "临近每日重置时间点（如23:59:50）",
     "①在重置前后各完成一次不同账号的首购\n②核对奖励",
     "1.daily奖励/首购奖励/购买奖励无重复无遗漏\n2.reward_state跨天重置正确",
     "", "", "", ""],
    ["TC-E07", "异常-GM辅助", "GM重置购买记录后首购判定正确", "P2",
     "GM可清档/重置特权卡数据",
     "①GM重置月卡buy_times=0\n②再次购买",
     "1.按buy_times==0判定为首购，重新发放首购奖励\n2.验证首购判定确实基于buy_times字段",
     "", "", "", "测试辅助用例，验证isFirstBuy判定逻辑"],
]

# ============ Sheet3 待确认内容 ============
pending = [
    [1, "终身卡首购文案不一致：代码配置first_desc为'购买额外可得神话仙友随机碎片*50'，需求图2为'首次购买额外可得神话仙友随机碎片'（无*50字样）",
     "以哪个为准？若以需求图为准需更新cfg_privilege_card[3].first_desc（客户端json+服务端lua两处）"],
    [2, "需求图数值与当前配置不一致：月卡图'购买即刻得160' vs 配置180元宝；终身卡图'1250' vs 配置1280元宝",
     "需求图为旧版示意图还是需调整配置？建议以cfg_privilege_card.buy_reward实际配置验收"],
    [3, "已购月卡玩家通关第40关时完全不弹任何卡宣传界面（当前逻辑），是否需要其它引导（如红点/跑马灯）",
     "确认交互预期，避免'需求理解偏差'"],
    [4, "月卡过期后重新登录是否需要再次主动弹出购买界面（当前onNewSysOpen仅在系统解锁事件触发一次，过期后不再弹）",
     "商业化为：确认过期召回是否在本需求范围内"],
    [5, "首购奖励在奖励弹窗中的展示顺序：当前服务端合并顺序为每日奖励→首购奖励→购买奖励",
     "确认弹窗内排序与图标数量上限（终身卡首购合计9种道具是否单页放得下）"],
    [6, "UIForeverCardOpenAlert界面及prefab是否需要彻底移除（当前仅注销触发逻辑，资源仍保留）",
     "确认保留（万一切回旧逻辑）或清理资源减小包体"],
]

# ============ 写Excel ============
wb = openpyxl.Workbook()
thin = Side(style="thin", color="BFBFBF")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
hdr_fill = PatternFill("solid", fgColor="4472C4")
hdr_font = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
cell_font = Font(name="微软雅黑", size=10)
wrap = Alignment(wrap_text=True, vertical="top", horizontal="left")
center = Alignment(wrap_text=True, vertical="center", horizontal="center")

def fill_sheet(ws, headers, rows, widths):
    ws.append(headers)
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = hdr_fill; cell.font = hdr_font; cell.alignment = center; cell.border = border
    for r in rows:
        ws.append(r)
    for r in range(2, ws.max_row + 1):
        for c in range(1, len(headers) + 1):
            cell = ws.cell(row=r, column=c)
            cell.font = cell_font; cell.alignment = wrap; cell.border = border
        # 优先级着色
        pc = ws.cell(row=r, column=4)
        if pc.value == "P0":
            pc.fill = PatternFill("solid", fgColor="F8CBAD"); pc.alignment = center
        elif pc.value == "P1":
            pc.fill = PatternFill("solid", fgColor="FFE699"); pc.alignment = center
        elif pc.value == "P2":
            pc.fill = PatternFill("solid", fgColor="C6E0B4"); pc.alignment = center
        # 测试结果下拉
        if len(headers) >= 10:
            from openpyxl.worksheet.datavalidation import DataValidation
    # 列宽/行高
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 22
    for r in range(2, ws.max_row + 1):
        ws.row_dimensions[r].height = None
    ws.freeze_panes = "A2"

# 数据验证下拉
from openpyxl.worksheet.datavalidation import DataValidation

ws1 = wb.active
ws1.title = "测试用例"
fill_sheet(ws1, HDR, cases, [9, 20, 30, 7, 26, 34, 44, 10, 20, 9, 30])
dv = DataValidation(type="list", formula1='"通过,失败,部分通过,阻塞,未执行,不适用"', allow_blank=True)
ws1.add_data_validation(dv)
dv.add(f"J2:J{ws1.max_row}")

ws2 = wb.create_sheet("异常与边界用例")
fill_sheet(ws2, HDR, exc_cases, [9, 16, 30, 7, 24, 34, 44, 10, 20, 9, 30])
dv2 = DataValidation(type="list", formula1='"通过,失败,部分通过,阻塞,未执行,不适用"', allow_blank=True)
ws2.add_data_validation(dv2)
dv2.add(f"J2:J{ws2.max_row}")

ws3 = wb.create_sheet("待确认内容")
fill_sheet(ws3, ["序号", "待确认内容", "说明/影响"], pending, [7, 60, 60])

wb.save(OUT)
print("saved:", OUT)
print("主流程用例:", len(cases), "异常用例:", len(exc_cases), "待确认:", len(pending))
