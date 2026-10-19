# -*- coding: utf-8 -*-
"""法宝第二轮复核结论回填（trae · 2026-09-09）——更新 法宝UI优化测试用例_trae.xlsx
依据：第二轮运行时取证 v3b(11:33)/v3c/v3d(12:14)/v3e(12:25)/v3f(12:37)/v3g(12:45)/v3h(12:51)
仅更新第二轮有新增证据的用例；其余沿用一轮结论。"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

DST = r"E:\workproject\temp_case\cases\法宝UI优化测试用例_trae.xlsx"
OUT = DST

# 第二轮复核：{TC: (新结论, 结果标签)}   结果标签→颜色见 STATUS
ROUND2 = {
    "TC007": ("未实现(第二轮运行时复现缺陷)：运行时patch 6002.ring_group 1→5后，ItemHelper.getItemQualityCfg(6002)外观数据源(atlas_res/color/bg/skill_res)完全不变=同quality不同group外观相同；图鉴分组键=e.quality(UISoulRingAtlas L132-137)；外观源只认quality不认ring_group，与静态D3一致", "未实现"),
    "TC008": ("未实现(第二轮运行时复现缺陷)：同group(1)下patch 6002.quality 1→7，qcfg_6002.atlas_res kuang07→kuang08、名称色随quality色变=同group不同quality外观不同；外观随quality变化", "未实现"),
    "TC009": ("未实现(第二轮运行时复现缺陷)：patch 6001.quality 1→7，qcfg_6001.atlas_res kuang07→kuang08=外观跟随quality变化；对照组patch quality[1].atlas_res→kuang03生效(v3h TC000-ctrl PASS)证明补丁链路与外观数据源=quality表；解耦未实现", "未实现"),
    "TC014": ("PASS(静态+第二轮运行时)：UISoulRingAtlasDesc L52『初始阶数』+formatLevel；D1旧chunk解决后图鉴详情实显『初始阶数：1阶』", "通过"),
    "TC020": ("PASS(第二轮运行时实锤，修正一轮误判)：图鉴可见项均显X阶(formatLevel)；全部年份残留(80万年等)均在【未拥有法宝的year标签】(RingAtlasItem L85仅在ringMap有该法宝时写入，否则保留预制体默认文本且activeNode隐藏不渲染)及active=False隐藏节点；v3b/c全UI扫描误判已修正", "通过"),
    "TC042": ("PASS(第二轮运行时)：GM award 6001+一键穿戴(BtnEquip_)后，装配面板+18/+18/+225与公式期望[18,18,225]完全一致(御剑符1阶addition=0.6%，3条属性全对)；角色面板SOUL_RING系统加成0.2%×3正常", "通过"),
    "TC043": ("PASS(第二轮运行时)：15字超长名项在视口内(|y|<400)，名称中心(379,306)位于图标中心(150,230)上方、无相交(相交0项)；截图人工复核通过", "通过"),
    "TC045": ("PASS(第二轮运行时)：D1解决后新chunk生效——VList虚拟滚动已工作(实例化项数20<全量72，随滚动回收)，分类ToggleContainer/TogContent链路正常，图鉴内容渲染正常", "通过"),
    "TC055": ("PASS(第二轮运行时)：炼宝页签UISoulRingDrawTabView切换打通——toggle走游戏setToggleList机制派发(node.emit('toggle'))，『法宝总览』→atlas激活=1、『炼宝』→draw激活=1；抽奖1次/10次扣费20→19→9精确(v3d)", "通过"),
    "TC044": ("部分实现(二轮复核加强)：分类仍按quality——运行时实锤图鉴datasMap按e.quality分组(UISoulRingAtlas L132-137)、列表按quality降序+已拥有置顶排序(L113-130)；筛选图数据化atlas_filter_res已引用；当前quality==ring_group故组数=10+全部", "部分实现"),
    "TC066": ("静态关注(二轮复核加强)：代码大量直接索引quality；第二轮每轮运行时仍实锤『cfg_soul_ring_quality索引-1』36~46条/轮、『cfg_soul_ring_pool索引0』1~2条/轮报错，另发现漫画页UIManHuaPage1.hideCallBack→SceneMap.playSceneEff空引用pageerror约3000-4600条/轮(与法宝无关，见D11)，建议一并补防御", "静态关注"),
}

STATUS = {
    "通过": ("通过", "FFC6EFCE"), "部分实现": ("部分实现", "FFFFEB9C"),
    "未实现": ("未实现", "FFFFC7CE"), "失败": ("失败", "FFFFC7CE"),
    "待运行时": ("待运行时(阻塞)", "FFD9D9D9"), "资源就绪": ("资源就绪", "FFFFEB9C"),
    "静态关注": ("静态关注", "FFFFEB9C"),
}

wb = openpyxl.load_workbook(DST)
ws = wb["测试用例"]

hdr_row = 1
id_col = 1
conc_col = None
res_col = None
for c in range(1, ws.max_column + 1):
    v = ws.cell(row=hdr_row, column=c).value
    if v == "用例编号":
        id_col = c
    if v and str(v).startswith("复测结论"):
        conc_col = c
    if v == "结果":
        res_col = c
assert conc_col and res_col, f"列未找到 conc={conc_col} res={res_col}"

n_upd = 0
for row in range(hdr_row + 1, ws.max_row + 1):
    cid = ws.cell(row=row, column=id_col).value
    if not cid:
        continue
    cid = str(cid).strip()
    if cid not in ROUND2:
        continue
    conc, label = ROUND2[cid]
    c1 = ws.cell(row=row, column=conc_col, value=conc)
    c1.alignment = Alignment(wrap_text=True, vertical="top")
    color = STATUS.get(label, ("", "FFFFFFFF"))[1]
    c2 = ws.cell(row=row, column=res_col, value=label)
    c2.fill = PatternFill("solid", fgColor=color)
    c2.alignment = Alignment(vertical="center", horizontal="center")
    n_upd += 1

try:
    wb.save(OUT)
except PermissionError:
    OUT = r"E:\workproject\temp_case\cases\法宝UI优化测试用例_trae_v2.xlsx"
    wb.save(OUT)
print("已更新:", OUT)
print("更新用例数:", n_upd)

# 统计最终分布
from collections import Counter
stat = Counter()
for row in range(hdr_row + 1, ws.max_row + 1):
    v = ws.cell(row=row, column=res_col).value
    if v:
        stat[str(v).strip()] += 1
total = sum(stat.values())
print("最终分布:", dict(stat))
print("总计:", total)
for k in ["通过", "部分实现", "未实现", "失败", "无法验证(需造数)", "待运行时(阻塞)", "资源就绪", "静态关注"]:
    if k in stat:
        print(f"  {k}: {stat[k]} ({stat[k]/total*100:.0f}%)")
